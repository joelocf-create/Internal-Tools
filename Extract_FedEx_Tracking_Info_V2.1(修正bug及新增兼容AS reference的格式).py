import re
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox

from openpyxl import load_workbook
from pypdf import PdfReader


APP_TITLE = "FedEx寄件報告轉檔輸出"
APP_VERSION = "V2.1"
TARGET_SHEET = "Tracking#"
START_ROW = 2
HEADERS = ("Master No.", "Package No.", "Weight", "EIP#")


def extract_fedex_data(pdf_path):
    """由 FedEx IPD/IED/IDF CLEARANCE REPORT 取出 Master No./Package No./Weight/Reference。"""
    reader = PdfReader(str(pdf_path))
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)

    if not full_text.strip():
        raise ValueError(f"{Path(pdf_path).name}：PDF 沒有可讀取的文字層。")

    # 每筆資料固定由 Master No. + Package No. + Weight 開始。
    # 先定位每筆起點，再於該筆範圍內抓取第一個 EPI 編號，
    # 可避免收件人姓名/地址換行造成 Regex 跨筆誤抓。
    start_pattern = re.compile(
        r"(?P<master>\d{12})\s+"
        r"(?P<package>\d{12})\s+"
        r"(?P<weight>\d+(?:\.\d+)?)"
    )
    starts = list(start_pattern.finditer(full_text))

    if not starts:
        raise ValueError(f"{Path(pdf_path).name}：未找到 FedEx 寄件明細資料。")

    records = []

    for i, match in enumerate(starts):
        record_end = starts[i + 1].start() if i + 1 < len(starts) else len(full_text)
        record_text = full_text[match.end():record_end]

        # Reference 擷取規則：
        # 1. 若該筆資料中有 EPI 編號，優先只取 EPI 資料。
        # 2. 若沒有 EPI，則取 "(CTN_xxxx)" 前方的完整 Reference 字串。
        eip_match = re.search(r"EPI\d+", record_text, re.IGNORECASE)

        if eip_match:
            # 有 EPI 時，只取 EPI 編號本身。
            reference_value = eip_match.group(0)
        else:
            # 無 EPI 時，取完整 Reference + CTN 字串，
            # 例如：415230235 (CTN_0043)
            #      CVAL30416006 (CTN_0054)
            #      15692032-1 (CTN_0001)
            #      360858853_739 (CTN_0003)
            reference_match = re.search(
                r"(?P<reference>CVAL[A-Za-z0-9_-]+|\d[A-Za-z0-9_-]*)"
                r"\s*(?P<ctn>\(CTN_[^)]+\))",
                record_text,
                re.IGNORECASE,
            )

            if not reference_match:
                raise ValueError(
                    f"{Path(pdf_path).name}：Package No. {match.group('package')} "
                    f"未找到 Reference，為避免輸出錯誤資料，程序已停止。"
                )

            reference_value = (
                f"{reference_match.group('reference')} "
                f"{reference_match.group('ctn')}"
            )

        records.append(
            (
                int(match.group("master")),
                int(match.group("package")),
                float(match.group("weight")),
                reference_value,
            )
        )

    return records



def write_to_excel(excel_path, records):
    """清除 Tracking# A2:D最後列舊資料，再寫入新資料並原檔存檔。"""
    excel_path = Path(excel_path)
    keep_vba = excel_path.suffix.lower() == ".xlsm"

    wb = load_workbook(excel_path, keep_vba=keep_vba)

    try:
        if TARGET_SHEET in wb.sheetnames:
            ws = wb[TARGET_SHEET]
        else:
            ws = wb.create_sheet(TARGET_SHEET)
            for col_index, header in enumerate(HEADERS, start=1):
                ws.cell(row=1, column=col_index, value=header)

        # 若既有工作表第1列抬頭有缺漏，只補空白抬頭，不覆蓋原有內容。
        for col_index, header in enumerate(HEADERS, start=1):
            if ws.cell(row=1, column=col_index).value in (None, ""):
                ws.cell(row=1, column=col_index, value=header)

        # 清除 A:D 第2列到最後有資料列，保留第1列抬頭與其他欄位。
        last_row = START_ROW - 1
        for col_index in range(1, 5):
            col_last_row = ws.max_row
            while (
                col_last_row >= START_ROW
                and ws.cell(row=col_last_row, column=col_index).value is None
            ):
                col_last_row -= 1
            last_row = max(last_row, col_last_row)

        if last_row >= START_ROW:
            for row in ws.iter_rows(
                min_row=START_ROW,
                max_row=last_row,
                min_col=1,
                max_col=4,
            ):
                for cell in row:
                    cell.value = None

        # 寫入 A:D。
        for row_index, (master_no, package_no, weight, eip_no) in enumerate(
            records,
            start=START_ROW,
        ):
            ws.cell(row=row_index, column=1, value=master_no)
            ws.cell(row=row_index, column=2, value=package_no)
            ws.cell(row=row_index, column=3, value=weight)
            ws.cell(row=row_index, column=4, value=eip_no)

            # Master No./Package No. 寫入為數字格式。
            ws.cell(row=row_index, column=1).number_format = "0"
            ws.cell(row=row_index, column=2).number_format = "0"
            ws.cell(row=row_index, column=3).number_format = "0.00"
            ws.cell(row=row_index, column=4).number_format = "@"

        wb.save(excel_path)
    finally:
        wb.close()


class FedExConverterApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"{APP_TITLE} {APP_VERSION}")
        self.root.resizable(False, False)

        self.pdf_paths = []
        self.excel_path = ""

        self.build_ui()

    def build_ui(self):
        main_frame = tk.Frame(self.root, padx=20, pady=18)
        main_frame.pack(fill="both", expand=True)

        title_label = tk.Label(
            main_frame,
            text=f"{APP_TITLE} {APP_VERSION}",
            font=("Microsoft JhengHei", 14, "bold"),
        )
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 16))

        pdf_button = tk.Button(
            main_frame,
            text="選擇FedEx報告",
            width=22,
            height=2,
            command=self.select_pdfs,
        )
        pdf_button.grid(row=1, column=0, padx=(0, 12), pady=6)

        excel_button = tk.Button(
            main_frame,
            text="選擇要輸出的Excel檔案",
            width=22,
            height=2,
            command=self.select_excel,
        )
        excel_button.grid(row=1, column=1, pady=6)

        self.pdf_status = tk.Label(
            main_frame,
            text="尚未選擇FedEx報告",
            anchor="w",
            justify="left",
            font=("Microsoft JhengHei", 9),
        )
        self.pdf_status.grid(row=2, column=0, columnspan=2, sticky="w", pady=(10, 4))

        self.excel_status = tk.Label(
            main_frame,
            text="尚未選擇Excel檔案",
            anchor="w",
            justify="left",
            font=("Microsoft JhengHei", 9),
        )
        self.excel_status.grid(row=3, column=0, columnspan=2, sticky="w", pady=4)

        self.run_button = tk.Button(
            main_frame,
            text="開始輸出",
            width=20,
            command=self.process_files,
            state="disabled",
        )
        self.run_button.grid(row=4, column=0, columnspan=2, pady=(16, 10))

        warning_label = tk.Label(
            main_frame,
            text="寄件資料會輸出至選定EXCEL檔的[Tracking#]工作表！",
            font=("Microsoft JhengHei", 9, "bold"),
            fg="red",
        )
        warning_label.grid(row=5, column=0, columnspan=2, pady=(4, 0))

    def select_pdfs(self):
        paths = filedialog.askopenfilenames(
            title="選擇 FedEx PDF Report",
            filetypes=[("PDF files", "*.pdf")],
        )

        if not paths:
            return

        self.pdf_paths = list(paths)
        self.pdf_status.config(text=f"已選擇 FedEx 報告：{len(self.pdf_paths)} 個檔案")
        self.update_run_button()

    def select_excel(self):
        path = filedialog.askopenfilename(
            title="選擇要寫入的 Excel 檔案",
            filetypes=[
                ("Excel files", "*.xlsx *.xlsm"),
                ("Excel Workbook", "*.xlsx"),
                ("Excel Macro-Enabled Workbook", "*.xlsm"),
            ],
        )

        if not path:
            return

        self.excel_path = path
        self.excel_status.config(text=f"輸出檔案：{Path(path).name}")
        self.update_run_button()

    def update_run_button(self):
        if self.pdf_paths and self.excel_path:
            self.run_button.config(state="normal")
        else:
            self.run_button.config(state="disabled")

    def process_files(self):
        try:
            self.run_button.config(state="disabled")
            self.root.config(cursor="wait")
            self.root.update_idletasks()

            all_records = []
            for pdf_path in self.pdf_paths:
                all_records.extend(extract_fedex_data(pdf_path))

            if not all_records:
                raise ValueError("未取得任何可輸出的 FedEx 寄件資料。")

            write_to_excel(self.excel_path, all_records)

            total_weight = sum(record[2] for record in all_records)
            messagebox.showinfo(
                "完成",
                f"FedEx 寄件資訊已匯入完成！\n\n"
                f"FedEx 報告：{len(self.pdf_paths):,} 個\n"
                f"資料筆數：{len(all_records):,}\n"
                f"重量合計：{total_weight:,.2f}\n"
                f"工作表：{TARGET_SHEET}",
            )

        except PermissionError:
            messagebox.showerror(
                "檔案無法寫入",
                "Excel 檔案可能正在開啟中，請先關閉 Excel 檔案後再執行。",
            )
        except Exception as exc:
            messagebox.showerror("執行錯誤", str(exc))
        finally:
            self.root.config(cursor="")
            self.update_run_button()


def main():
    root = tk.Tk()
    FedExConverterApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
