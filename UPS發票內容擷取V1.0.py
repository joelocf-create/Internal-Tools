"""
UPS電子發票PDF擷取工具  '2026/07/29 10:30
- 選擇一個或多個PDF檔案（每頁為一張發票）
- 擷取：發票日期/發票號碼/品名/帳單號碼/銷售額合計/營業稅/總計
- 輸出Excel存於來源檔案同資料夾，檔名為 發票擷取結果_YYYYMMDD_HHMMSS.xlsx
- 無法擷取完整資料的頁面，該列反黃標記
- 處理過程中於Terminal顯示逐頁進度
"""

import re
import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

# ── 欄位擷取用正則 ──────────────────────────────
RE_DATE = re.compile(r"(\d{4})\s*-\s*(\d{1,2})\s*-\s*(\d{1,2})")
RE_INVOICE_NO = re.compile(r"發票號碼[:：]\s*([A-Za-z0-9]+)")
RE_ITEM = re.compile(r"品名\s*數量\s*單價\s*金額\s*備註\s*\n(\S+)")
RE_BILL_NO = re.compile(r"帳單號碼[：:]\s*\n?(\d+)")
RE_SALES_TOTAL = re.compile(r"銷售額合計\s+([\d,]+)")
RE_TAX = re.compile(r"應稅\s+V\s+零稅率\s+免稅\s+([\d,]+)")
RE_TOTAL = re.compile(r"總計\s+([\d,]+)")

FIELDS = ["發票日期", "發票號碼", "品名", "帳單號碼", "銷售額合計", "營業稅", "總計"]

HIGHLIGHT_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def extract_page(text: str) -> dict:
    """從單頁純文字中擷取所需欄位，擷取失敗的欄位值為None"""
    result = {}

    m = RE_DATE.search(text)
    if m:
        year, month, day = map(int, m.groups())
        try:
            result["發票日期"] = datetime.date(year, month, day)
        except ValueError:
            result["發票日期"] = None
    else:
        result["發票日期"] = None

    m = RE_INVOICE_NO.search(text)
    result["發票號碼"] = m.group(1) if m else None

    m = RE_ITEM.search(text)
    result["品名"] = m.group(1) if m else None

    m = RE_BILL_NO.search(text)
    result["帳單號碼"] = m.group(1) if m else None

    m = RE_SALES_TOTAL.search(text)
    result["銷售額合計"] = int(m.group(1).replace(",", "")) if m else None

    m = RE_TAX.search(text)
    result["營業稅"] = int(m.group(1).replace(",", "")) if m else None

    m = RE_TOTAL.search(text)
    result["總計"] = int(m.group(1).replace(",", "")) if m else None

    return result


def process_files(file_paths: list[str]) -> list[dict]:
    """處理多個PDF檔案，回傳每頁的資料列（含來源檔名與頁碼），並於Terminal顯示進度"""
    rows = []
    total_files = len(file_paths)
    for file_idx, file_path in enumerate(file_paths, start=1):
        file_name = Path(file_path).name
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            print(f"[{file_idx}/{total_files}] 處理檔案：{file_name}（共 {total_pages} 頁）")
            for page_index, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                data = extract_page(text)
                data["來源檔案名稱"] = file_name
                data["頁碼"] = page_index
                rows.append(data)
                status = "OK" if all(data[field] is not None for field in FIELDS) else "缺漏"
                print(f"    第 {page_index}/{total_pages} 頁擷取完成 [{status}]")
    print(f"全部處理完成，共 {len(rows)} 頁。")
    return rows


def build_excel(rows: list[dict], output_path: str) -> None:
    """將擷取結果寫入Excel，缺漏欄位的列反黃標記，並自動調整欄寬與字型"""
    wb = Workbook()
    ws = wb.active
    ws.title = "發票資料"

    headers = ["來源檔案名稱", "頁碼"] + FIELDS
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    for data in rows:
        row_values = [data["來源檔案名稱"], data["頁碼"]] + [data[field] for field in FIELDS]
        ws.append(row_values)
        row_idx = ws.max_row

        # 日期欄格式化為正常日期顯示
        date_cell = ws.cell(row=row_idx, column=headers.index("發票日期") + 1)
        if isinstance(date_cell.value, datetime.date):
            date_cell.number_format = "yyyy/mm/dd"

        # 金額欄加千分位格式
        for field in ("銷售額合計", "營業稅", "總計"):
            amount_cell = ws.cell(row=row_idx, column=headers.index(field) + 1)
            if amount_cell.value is not None:
                amount_cell.number_format = "#,##0"

        # 任一必要欄位缺漏 → 整列反黃
        row_missing = any(data[field] is None for field in FIELDS)
        if row_missing:
            for cell in ws[row_idx]:
                cell.fill = HIGHLIGHT_FILL

        for cell in ws[row_idx]:
            cell.font = Font(name="Arial")

    # 自動調整欄寬
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_length = max(
            (len(str(cell.value)) for cell in ws[col_letter] if cell.value is not None),
            default=8,
        )
        ws.column_dimensions[col_letter].width = max_length + 4

    wb.save(output_path)


def main() -> None:
    root = tk.Tk()
    root.withdraw()

    file_paths = filedialog.askopenfilenames(
        title="請選擇要處理的發票PDF檔案（可多選）",
        filetypes=[("PDF files", "*.pdf")],
    )
    if not file_paths:
        return

    print(f"已選取 {len(file_paths)} 個檔案，開始處理...")
    rows = process_files(list(file_paths))
    if not rows:
        messagebox.showwarning("提示", "所選檔案中未擷取到任何頁面資料。")
        return

    # 輸出Excel存於第一個來源檔案同資料夾
    output_dir = Path(file_paths[0]).parent
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"發票擷取結果_{timestamp}.xlsx"

    build_excel(rows, str(output_path))

    missing_count = sum(
        1 for data in rows if any(data[field] is None for field in FIELDS)
    )
    msg = f"完成，共處理 {len(rows)} 頁，輸出至：\n{output_path}"
    if missing_count:
        msg += f"\n\n有 {missing_count} 列資料擷取不完整，已於Excel中反黃標記。"
    messagebox.showinfo("完成", msg)


if __name__ == "__main__":
    main()
