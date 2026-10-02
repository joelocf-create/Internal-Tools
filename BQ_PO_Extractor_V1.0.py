import re
import pymupdf
from datetime import datetime
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# B&Q (Kingfisher) PO PDF Extractor
# 2026/10/02
#
# 功能簡述：
# 1. 處理 B&Q Limited Purchase Order PDF（Article to Ship + Shipment Schedule）
# 2. 擷取：PO#、Banner Article No.、Article Description、Quantity、
#          Unit Cost、Amount、POD、LRD
# 3. 以 Shipment Schedule 為主（同一品項分多個 Shipment 時會各出一列）；
#    Article Description 取自 Article to Ship
# 4. 可多檔案批次，Excel 可自選路徑
# 5. 驗證：Quantity × Unit Cost = Amount；Amount 合計 = PDF Total
# ============================================================


SHEET_NAME = "PO_Data"
HEADERS = [
    "PO#", "Banner Article No.", "Article Description",
    "Quantity", "Unit Cost", "Amount", "POD", "LRD",
]
AMOUNT_TOLERANCE = 0.011  # 數量 × 單價 與金額的允許誤差

RE_KAL_ART = re.compile(r"^PC\d{8}$")
RE_BANNER_ART = re.compile(r"^\d{13}$")
RE_DATE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
RE_MONEY = re.compile(r"^[\d,]+\.\d{2}$")


def to_number(text: str) -> float:
    """'1,228' / '4,365.00' -> 數值"""
    return float(text.replace(",", ""))


def read_pdf(pdf_path: str):
    """讀取 PDF，回傳 (全文文字, 每頁 words 清單)"""
    texts = []
    pages_words = []

    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            texts.append(page.get_text("text"))
            # 此 PDF 頁面為旋轉頁（rotation=90），words 座標需轉成視覺座標，
            # 否則欄位 X / 列 Y 方向會互換
            matrix = page.rotation_matrix
            visual = []

            for w in page.get_text("words"):
                r = pymupdf.Rect(w[:4]) * matrix
                r.normalize()
                visual.append((r.x0, r.y0, r.x1, r.y1, w[4]))

            pages_words.append(visual)

    return "\n".join(texts), pages_words


def extract_po_no(text: str) -> str:
    """抓 PO#，範例：Purchase  Order: 137284049"""
    m = re.search(r"Purchase\s+Order:\s*(\d+)", text)
    return m.group(1) if m else ""


def extract_articles(text: str) -> dict:
    """
    功能：
    擷取 Article to Ship 區段，回傳 {KAL Art No: {...}}。

    重點：
    1. 每頁重複的欄位標題先移除，避免跨頁品項被切斷或誤判 '|'。
    2. 以 PCxxxxxxxx 作為每筆品項起點。
    3. Description 為 Banner Article No.(13 碼) 之後、以 '|' 結尾的文字。
    """
    # 只取 Article to Ship 區段（Shipment Schedule 之前）
    cut = text.find("Shipment Schedule")
    article_text = text[:cut] if cut >= 0 else text

    # 移除每頁重複標題 / 頁尾
    article_text = re.sub(
        r"Article to Ship\n(?:.*\n)*?Outer Bar Code\n", "", article_text
    )
    article_text = re.sub(r"(?m)^Page \d+ of \d+\s*$", "", article_text)
    article_text = re.sub(r"(?m)^Last Update Date:.*$", "", article_text)

    # 截掉 Total: 之後內容
    total_pos = article_text.find("Total:")
    if total_pos >= 0:
        article_text = article_text[:total_pos]

    parts = re.split(r"(?m)^(PC\d{8})\s*$", article_text)
    articles = {}

    # parts = [前置文字, PC1, 區塊1, PC2, 區塊2, ...]
    for i in range(1, len(parts) - 1, 2):
        kal_no = parts[i]
        lines = [x.strip() for x in parts[i + 1].splitlines() if x.strip()]

        # lines: Vendor Art / Factory / Qty / Unit / Unit Cost / ...
        if len(lines) < 5:
            continue

        try:
            qty = int(to_number(lines[2]))
            unit_cost = to_number(lines[4])
        except ValueError:
            continue

        amount = None
        for line in lines[5:]:
            if RE_MONEY.match(line):
                amount = to_number(line)
                break

        # Description：第一個含 '|' 的行往回找到 13 碼 Banner Article No.
        banner_no = ""
        description = ""
        pipe_idx = next((k for k, x in enumerate(lines) if "|" in x), None)

        if pipe_idx is not None:
            start = None
            for k in range(pipe_idx, -1, -1):
                if RE_BANNER_ART.match(lines[k]):
                    start = k
                    break

            if start is not None:
                banner_no = lines[start]
                desc_lines = lines[start + 1:pipe_idx + 1]
                description = " ".join(desc_lines).split("|")[0]
                description = re.sub(r"\s+", " ", description).strip()

        articles[kal_no] = {
            "Banner Article No.": banner_no,
            "Article Description": description,
            "Quantity": qty,
            "Unit Cost": unit_cost,
            "Amount": amount,
        }

    return articles


def extract_shipments(pages_words: list) -> list[dict]:
    """
    功能：
    以座標法擷取 Shipment Schedule 每一列。

    POD 在 PDF 內會被拆成多行（UNITED / KINGDOM / FELIXSTOWE / (SAP Easier)），
    文字模式無法分辨欄位，因此以標題 POD ~ DC 的 X 座標範圍取字。
    """
    shipments = []

    for words in pages_words:
        # words: (x0, y0, x1, y1, text)，已轉為視覺座標
        header_pod = next((w for w in words if w[4] == "POD"), None)
        header_dc = next((w for w in words if w[4] == "DC"), None)
        header_lrd = next((w for w in words if w[4] == "LRD"), None)

        if not (header_pod and header_dc and header_lrd):
            continue

        pod_x0 = header_pod[0] - 3
        dc_x0 = header_dc[0] - 3
        lrd_x0 = header_lrd[0] - 3

        # 列錨點：LRD 欄右側的 KAL Art No (PCxxxxxxxx)
        anchors = sorted(
            [w for w in words
             if RE_KAL_ART.match(w[4]) and w[0] > header_lrd[0]],
            key=lambda w: w[1]
        )

        if not anchors:
            continue

        # 區塊結束位置：頁面最下方或 "POL =" 說明列
        pol_note = next((w for w in words if w[4] == "POL"
                         and w[0] < header_pod[0] and w[1] > anchors[-1][1]
                         and w[1] > header_lrd[1]), None)
        page_end_y = pol_note[1] if pol_note else 1e9

        # 列範圍：錨點本身 y 為每列第一行（Shipment No. / SEA 同行）
        for idx, anchor in enumerate(anchors):
            y_top = anchor[1] - 2
            y_bottom = anchors[idx + 1][1] - 2 if idx + 1 < len(anchors) \
                else page_end_y

            row_words = [w for w in words if y_top <= w[1] < y_bottom]

            # --- POD ---
            pod_words = sorted(
                [w for w in row_words if pod_x0 <= w[0] < dc_x0],
                key=lambda w: (round(w[1]), w[0])
            )
            pod = " ".join(w[4] for w in pod_words)

            # --- LRD：同一行、KAL Art No 左側第一個日期 ---
            same_line = sorted(
                [w for w in row_words if abs(w[1] - anchor[1]) < 3],
                key=lambda w: w[0]
            )
            dates = [w[4] for w in same_line
                     if RE_DATE.match(w[4]) and w[0] >= lrd_x0]
            lrd = datetime.strptime(dates[0], "%d/%m/%Y") if dates else None

            # --- Qty / Unit / Unit Cost / Amount：錨點右側同一行 ---
            right = [w[4] for w in same_line if w[0] > anchor[0]]
            # 預期：Qty, Unit, Unit Cost, Ship Amount, Ship CBM
            if len(right) < 4:
                continue

            try:
                qty = int(to_number(right[0]))
                unit_cost = to_number(right[2])
                amount = to_number(right[3])
            except ValueError:
                continue

            shipments.append({
                "KAL Art No": anchor[4],
                "Quantity": qty,
                "Unit Cost": unit_cost,
                "Amount": amount,
                "POD": pod,
                "LRD": lrd,
            })

    return shipments


def extract_total_amount(text: str):
    """抓 Article to Ship 的 Total 金額（用於合計驗證）"""
    m = re.search(r"Total:\s*\n\s*([\d,]+)\s*\n\s*([\d,]+\.\d{2})", text)
    return to_number(m.group(2)) if m else None


def process_pdf(pdf_path: str) -> tuple[list[dict], list[str]]:
    """處理單一 PDF，回傳 (資料列, 警告訊息)"""
    text, pages_words = read_pdf(pdf_path)

    po_no = extract_po_no(text)
    articles = extract_articles(text)
    shipments = extract_shipments(pages_words)
    warnings = []

    if not articles:
        raise ValueError("找不到 Article to Ship 資料")

    rows = []

    if shipments:
        # 依 Article to Ship 的順序輸出（Shipment Schedule 排序與其不同）
        art_order = {k: i for i, k in enumerate(articles)}
        shipments.sort(key=lambda x: art_order.get(x["KAL Art No"], len(art_order)))

        for s in shipments:
            art = articles.get(s["KAL Art No"])

            if art is None:
                warnings.append(f"{s['KAL Art No']} 不在 Article to Ship 內")
                art = {"Banner Article No.": "", "Article Description": ""}

            rows.append({
                "PO#": po_no,
                "Banner Article No.": art["Banner Article No."],
                "Article Description": art["Article Description"],
                "Quantity": s["Quantity"],
                "Unit Cost": s["Unit Cost"],
                "Amount": s["Amount"],
                "POD": s["POD"],
                "LRD": s["LRD"],
            })
    else:
        # 無 Shipment Schedule 時，退而使用 Article to Ship（POD / LRD 空白）
        warnings.append("找不到 Shipment Schedule，POD / LRD 留白")
        for art in articles.values():
            rows.append({
                "PO#": po_no,
                "Banner Article No.": art["Banner Article No."],
                "Article Description": art["Article Description"],
                "Quantity": art["Quantity"],
                "Unit Cost": art["Unit Cost"],
                "Amount": art["Amount"],
                "POD": "",
                "LRD": None,
            })

    # ===== 驗證 =====
    for r in rows:
        if r["Amount"] is None or abs(
            r["Quantity"] * r["Unit Cost"] - r["Amount"]
        ) >= AMOUNT_TOLERANCE:
            warnings.append(
                f"{r['Banner Article No.']} Qty×Unit Cost≠Amount"
            )

        if not r["Banner Article No."] or not r["Article Description"]:
            warnings.append(f"{r['Banner Article No.'] or '(無編號)'} 缺少編號或描述")

    pdf_total = extract_total_amount(text)
    row_total = round(sum(r["Amount"] or 0 for r in rows), 2)

    if pdf_total is not None and abs(row_total - pdf_total) >= AMOUNT_TOLERANCE:
        warnings.append(f"Amount 合計 {row_total:,.2f} ≠ PDF Total {pdf_total:,.2f}")

    return rows, warnings


def export_to_excel(all_rows: list[dict], output_path: str):
    """輸出 Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append(HEADERS)

    for r in all_rows:
        ws.append([r[h] for h in HEADERS])

    # ===== 格式設定 =====
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)

    col = {h: i + 1 for i, h in enumerate(HEADERS)}

    for row in range(2, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row, c).font = Font(name="Arial")

        ws.cell(row, col["PO#"]).number_format = "@"
        ws.cell(row, col["Banner Article No."]).number_format = "@"
        ws.cell(row, col["Quantity"]).number_format = "#,##0"
        ws.cell(row, col["Unit Cost"]).number_format = "#,##0.00"
        ws.cell(row, col["Amount"]).number_format = "#,##0.00"
        ws.cell(row, col["LRD"]).number_format = "yyyy/m/d"
        ws.cell(row, col["LRD"]).alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A2"

    # ===== 自動調整欄寬 =====
    for c in range(1, ws.max_column + 1):
        letter = get_column_letter(c)
        max_len = max(
            (len(str(cell.value)) for cell in ws[letter] if cell.value is not None),
            default=0
        )
        ws.column_dimensions[letter].width = min(max_len + 2, 60)

    wb.save(output_path)


def main():
    root = Tk()
    root.withdraw()

    # ===== 選 PDF（可多選）=====
    pdf_paths = filedialog.askopenfilenames(
        title="選擇 B&Q PO PDF",
        filetypes=[("PDF Files", "*.pdf")]
    )

    if not pdf_paths:
        return

    # ===== 選輸出位置 =====
    output_path = filedialog.asksaveasfilename(
        title="儲存 Excel",
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        initialfile="BQ_PO_Extract.xlsx"
    )

    if not output_path:
        return

    all_rows = []
    messages = []
    ok_count = 0

    # 逐檔處理，單檔失敗不中斷其他檔案
    for pdf in pdf_paths:
        name = Path(pdf).name

        try:
            rows, warnings = process_pdf(pdf)
            all_rows.extend(rows)
            ok_count += 1

            for w in warnings:
                messages.append(f"[警告] {name}：{w}")

        except Exception as e:
            messages.append(f"[失敗] {name}：{e}")

    if not all_rows:
        messagebox.showerror("失敗", "沒有可輸出的資料：\n" + "\n".join(messages))
        return

    try:
        export_to_excel(all_rows, output_path)
    except Exception as e:
        messagebox.showerror("失敗", f"Excel 寫入失敗（檔案是否已開啟？）：\n{e}")
        return

    summary = (
        f"完成輸出：{output_path}\n"
        f"成功 {ok_count}/{len(pdf_paths)} 檔，共 {len(all_rows)} 列"
    )

    if messages:
        summary += "\n\n" + "\n".join(messages)
        messagebox.showwarning("完成（含警告）", summary)
    else:
        messagebox.showinfo("完成", summary)


if __name__ == "__main__":
    main()
