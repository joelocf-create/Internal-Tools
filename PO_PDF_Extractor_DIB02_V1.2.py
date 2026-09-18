import re
import fitz
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# DIB02 PO PDF Extractor
# 2026/08/18
#
# 功能簡述：
# 1. 專門處理 DIB02 (TV Hardware Dist LLC) 格式
# 2. 擷取：
#    - PO#
#    - Ship When
#    - 表身資料
# 3. 可多檔案批次
# 4. Excel 可自選路徑
# 5. 檔名前綴 DIB02_
# 6. Item / Quantity / Unit Cost / Ext Cost 為必要資料；其餘欄位可空白
# 7. 支援產品欄位被 PDF 文字層拆成多列，以及 UPC 前導 0 被省略
# ============================================================


def extract_pdf_text(pdf_path: str) -> str:
    """讀取 PDF 文字"""
    text_list = []

    with fitz.open(pdf_path) as doc:
        for page in doc:
            text_list.append(page.get_text("text"))

    return "\n".join(text_list)


def clean_text(text: str) -> str:
    """清理文字"""
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def extract_po_no(text: str) -> str:
    """
    功能：
    抓 PO#

    範例：
    ... bills of lading -01292360W2900
    """

    m = re.search(r"-([A-Z0-9]+)", text)

    return m.group(1) if m else ""


def extract_ship_date(text: str) -> str:
    """
    功能：
    抓 SHIP WHEN

    範例：
    SHIP WHEN: 05/18/26 → 2026/05/18
    """

    m = re.search(r"SHIP WHEN:\s*(\d{2})/(\d{2})/(\d{2})", text)

    if not m:
        return ""

    mm, dd, yy = m.groups()
    yyyy = "20" + yy

    return f"{yyyy}/{mm}/{dd}"


def extract_product_lines(text: str) -> list[dict]:
    """
    功能：
    擷取 DIB02 表身資料。

    修正原則：
    1. Item / Quantity / Unit Cost / Ext Cost 為必要資料。
    2. UPC Code / MFG SKU / Description / Master Pack 若 PDF 無資料，
       則輸出空白。
    3. 不依賴產品資料是否在同一列；以 6 碼 Item 切分每筆產品區塊。
    4. UPC 允許 11~14 碼，支援 PDF 文字層省略前導 0。
    """

    products = []

    # 每筆產品以獨立一列的 6 碼 Item 作為起點。
    item_pattern = re.compile(r"(?m)^(?P<item>\d{6})\s*$")
    item_matches = list(item_pattern.finditer(text))

    for index, item_match in enumerate(item_matches):
        item = item_match.group("item")

        if index + 1 < len(item_matches):
            block_end = item_matches[index + 1].start()
        else:
            remaining_text = text[item_match.end():]
            total_match = re.search(
                r"(?mi)^Total\s+Items\s*:",
                remaining_text
            )

            if total_match:
                block_end = item_match.end() + total_match.start()
            else:
                block_end = len(text)

        block_text = text[item_match.end():block_end]

        # ===== 1. Unit Cost / Ext Cost =====
        # DIB02 價格欄有 $，以此作為最穩定定位。
        price_matches = re.findall(
            r"\$\s*([\d,]+(?:\.\d+)?)",
            block_text
        )

        if len(price_matches) < 2:
            continue

        unit_cost = float(price_matches[0].replace(",", ""))
        ext_cost = float(price_matches[1].replace(",", ""))

        # ===== 2. Quantity / UPC =====
        # 實際 PDF 文字層常將 CURR QTY 與 UPC 放在同一列。
        # UPC 有時為 14 碼，有時前導 0 被移除只剩 11 碼。
        qty_upc_match = re.search(
            r"(?m)^\s*(\d[\d,]*)\s+(\d{11,14})\s*$",
            block_text
        )

        quantity = None
        upc = ""

        if qty_upc_match:
            quantity = int(qty_upc_match.group(1).replace(",", ""))
            upc = qty_upc_match.group(2)

        # 若 UPC 整欄缺值，仍允許解析。
        # 對區塊內的純整數逐一測試，以 Qty × Unit Cost = Ext Cost 確認 Quantity。
        if quantity is None:
            qty_candidates = re.findall(
                r"(?m)^\s*(\d[\d,]*)\s*$",
                block_text
            )

            for token in qty_candidates:
                candidate_qty = int(token.replace(",", ""))

                if abs(candidate_qty * unit_cost - ext_cost) < 0.011:
                    quantity = candidate_qty
                    break

        if quantity is None:
            continue

        # 必要欄位最終驗證。
        if abs(quantity * unit_cost - ext_cost) >= 0.011:
            continue

        # ===== 3. MFG SKU（optional）=====
        sku = ""

        if qty_upc_match:
            after_upc = block_text[qty_upc_match.end():]

            sku_match = re.search(
                r"(?m)^\s*([A-Z0-9]*[A-Z][A-Z0-9]*)\s*$",
                after_upc
            )

            if sku_match:
                sku = sku_match.group(1)

        # ===== 4. Description（optional）=====
        description = ""

        if sku:
            sku_line_match = re.search(
                rf"(?m)^\s*{re.escape(sku)}\s*$",
                block_text
            )

            if sku_line_match:
                after_sku = block_text[sku_line_match.end():]
                first_price_pos = after_sku.find("$")

                if first_price_pos >= 0:
                    desc_area = after_sku[:first_price_pos]
                else:
                    desc_area = after_sku

                desc_lines = [
                    re.sub(r"\s+", " ", x).strip()
                    for x in desc_area.splitlines()
                    if x.strip()
                ]

                # 價格前最後一個純整數通常是 Master Pack。
                if desc_lines and re.fullmatch(r"\d+", desc_lines[-1]):
                    desc_lines.pop()

                if desc_lines:
                    description = " ".join(desc_lines).strip()

        # ===== 5. Master Pack（optional）=====
        master_pack = ""

        first_price_pos = block_text.find("$")

        if first_price_pos >= 0:
            before_price = block_text[:first_price_pos]
            integer_lines = re.findall(
                r"(?m)^\s*(\d+)\s*$",
                before_price
            )

            # 排除 ORG QTY；最靠近價格的純整數通常為 Pack。
            if integer_lines:
                candidate_pack = int(integer_lines[-1])

                if candidate_pack != quantity:
                    master_pack = candidate_pack
                elif len(integer_lines) >= 2:
                    # 若最後一個剛好等於 Quantity，再檢查前一個。
                    previous_pack = int(integer_lines[-2])

                    if previous_pack != quantity:
                        master_pack = previous_pack

        products.append({
            "Item": item,
            "Description": description,
            "MFG SKU": sku,
            "Quantity": quantity,
            "UPC Code": upc,
            "Master Pack": master_pack,
            "Unit Cost": unit_cost,
            "Ext Cost": ext_cost,
        })

    return products


def export_to_excel(all_rows: list[dict], output_path: str):
    """輸出 Excel"""

    wb = Workbook()
    ws = wb.active
    ws.title = "PO_Data"

    headers = [
        "File Name", "PO#", "Ship Date",
        "Item", "Description", "MFG SKU",
        "Quantity", "UPC Code",
        "Master Pack", "Unit Cost", "Ext Cost"
    ]

    ws.append(headers)

    for row in all_rows:
        ws.append([
            row["File Name"],
            row["PO#"],
            row["Ship Date"],
            row["Item"],
            row["Description"],
            row["MFG SKU"],
            row["Quantity"],
            row["UPC Code"],
            row["Master Pack"],
            row["Unit Cost"],
            row["Ext Cost"],
        ])

    # ===== 格式設定 =====
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    # ===== 數字格式設定 =====
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 8).number_format = "@"          # UPC Code 保留前導 0
        ws.cell(row, 10).number_format = "#,##0.000" # Unit Cost，小數點三位
        ws.cell(row, 11).number_format = "#,##0.00"  # Ext Cost，小數點二位

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ===== 自動調整欄寬 =====
    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        max_len = 0

        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    wb.save(output_path)


def main():
    root = Tk()
    root.withdraw()

    # ===== 選 PDF =====
    pdf_paths = filedialog.askopenfilenames(
        title="選擇 DIB02 PDF",
        filetypes=[("PDF Files", "*.pdf")]
    )

    if not pdf_paths:
        return

    # ===== 選輸出位置 =====
    output_path = filedialog.asksaveasfilename(
        title="儲存 Excel",
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        initialfile="DIB02_PO_Extract.xlsx"
    )

    if not output_path:
        return

    all_rows = []

    for pdf in pdf_paths:
        text = clean_text(extract_pdf_text(pdf))

        po = extract_po_no(text)
        ship = extract_ship_date(text)
        products = extract_product_lines(text)

        for p in products:
            p["File Name"] = Path(pdf).name
            p["PO#"] = po
            p["Ship Date"] = ship
            all_rows.append(p)

    export_to_excel(all_rows, output_path)

    messagebox.showinfo("完成", f"完成輸出：\n{output_path}")


if __name__ == "__main__":
    main()