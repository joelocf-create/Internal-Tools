from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

import fitz  # PyMuPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


@dataclass
class OrderItem:
    po_number: str
    kal_article_no: str
    banner_article_no: str
    article_description: str
    quantity: int
    unit_cost: float
    amount: float
    pod: str = ""
    lrd: str = ""


@dataclass
class ValidationResult:
    po_number: str
    pdf_total_qty: int
    pdf_total_amount: float
    output_total_qty: int
    output_total_amount: float
    is_match: bool


HEADER = [
    "PO#",
    "Banner Article No.",
    "Article Description",
    "Quantity",
    "Unit Cost",
    "Amount",
    "POD",
    "LRD",
]


def clean_lines(text: str) -> List[str]:
    """將 PDF 文字切成乾淨的非空白列。"""
    return [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]


def extract_po_number(full_text: str) -> str:
    """擷取 Purchase Order 或 Proforma Order 編號。"""
    match = re.search(
        r"(?:Purchase|Proforma)\s+Order:\s*(\d+)",
        full_text,
        re.IGNORECASE,
    )
    if not match:
        raise ValueError("找不到 Purchase Order / Proforma Order 編號。")
    return match.group(1)


def extract_pdf_total(full_text: str) -> tuple[int, float]:
    """擷取 Article to Ship 區段的 Total 數量與總金額。"""
    article_section = full_text.split("Shipment Schedule", 1)[0]

    matches = re.findall(
        r"Total:\s*([\d,]+)\s+([\d,]+\.\d{2})",
        article_section,
        re.IGNORECASE,
    )
    if not matches:
        raise ValueError("找不到訂單 Total 的總數量及總金額。")

    qty_text, amount_text = matches[-1]
    return int(qty_text.replace(",", "")), float(amount_text.replace(",", ""))


def extract_article_items(doc: fitz.Document, po_number: str) -> List[OrderItem]:
    """解析 Article to Ship 區段。"""
    article_text_parts: List[str] = []

    for page in doc:
        text = page.get_text("text")
        if "Article to Ship" in text:
            if "Shipment Schedule" in text:
                text = text.split("Shipment Schedule", 1)[0]
            article_text_parts.append(text)

    if not article_text_parts:
        raise ValueError("找不到 Article to Ship 區段。")

    lines = clean_lines("\n".join(article_text_parts))

    # 找出每個 KAL Article 的起始位置。此格式的料號為 PC + 數字。
    starts = [i for i, line in enumerate(lines) if re.fullmatch(r"PC\d+", line)]
    if not starts:
        raise ValueError("找不到 Article 明細。")

    items: List[OrderItem] = []

    for idx, start in enumerate(starts):
        end = starts[idx + 1] if idx + 1 < len(starts) else len(lines)
        block = lines[start:end]

        try:
            kal_article_no = block[0]
            quantity = int(block[3].replace(",", ""))
            unit_cost = float(block[5].replace(",", ""))

            barcode_indexes = [
                i for i, value in enumerate(block)
                if re.fullmatch(r"\d{12,14}", value)
            ]
            if not barcode_indexes:
                raise ValueError("找不到 Banner Article No.")

            banner_article_no = block[barcode_indexes[0]]

            # Amount 不再依賴固定的 Banner Article Cat Desc. 文字。
            # 不同國家 PO 的分類描述可能不同，例如：
            # Poland: "17 Window & Textiles"
            # France: "CAS-3 DECO == 47 DECO / INT == 572 / STORE**STORE"
            # 因此改用 Qty × Unit Cost 的應付金額，從第一個與第二個條碼之間
            # 的數值中尋找相符的 Amount，以提高不同版型的相容性。
            amount: Optional[float] = None
            expected_amount = round(quantity * unit_cost, 2)

            amount_search_start = barcode_indexes[0] + 1
            amount_search_end = barcode_indexes[1] if len(barcode_indexes) > 1 else len(block)

            for value in block[amount_search_start:amount_search_end]:
                if not re.fullmatch(r"[\d,]+\.\d{2}", value):
                    continue

                candidate = float(value.replace(",", ""))
                if abs(candidate - expected_amount) < 0.005:
                    amount = candidate
                    break

            if amount is None:
                raise ValueError(
                    f"找不到 Amount（預期約為 {expected_amount:,.2f}）。"
                )

            # Product title 在第二次出現的 Banner Article No. 後第一列。
            # 不跨頁追 Ranged，避免把下一頁頁首/欄位抬頭誤併入品名。
            desc_start = barcode_indexes[1] + 1 if len(barcode_indexes) > 1 else barcode_indexes[0] + 1
            article_description = block[desc_start].strip()
            article_description = re.sub(r"\s*\|\s*Ranged\s*$", "", article_description, flags=re.IGNORECASE).strip()
            article_description = re.sub(r"\s*\|\s*$", "", article_description).strip()
            article_description = re.sub(r"\s+", " ", article_description)

            if not article_description:
                raise ValueError("找不到 Article Description。")

            items.append(
                OrderItem(
                    po_number=po_number,
                    kal_article_no=kal_article_no,
                    banner_article_no=banner_article_no,
                    article_description=article_description,
                    quantity=quantity,
                    unit_cost=unit_cost,
                    amount=amount,
                )
            )
        except (IndexError, ValueError) as exc:
            raise ValueError(f"解析 {block[0] if block else '未知項目'} 失敗：{exc}") from exc

    return items


def extract_shipment_info(doc: fitz.Document) -> Dict[str, Dict[str, str]]:
    """依 KAL Article No. 解析 POD 與 LRD。"""
    shipment_map: Dict[str, Dict[str, str]] = {}

    for page in doc:
        text = page.get_text("text")
        if "Shipment Schedule" not in text:
            continue

        lines = clean_lines(text)
        for i, line in enumerate(lines):
            # PDF 可能把兩個日期與 KAL Article 放在同一列。
            m = re.fullmatch(r"(\d{2}/\d{2}/\d{4})\s+(\d{2}/\d{2}/\d{4})\s+(PC\d+)", line)
            if not m:
                continue

            lrd = m.group(1)
            kal_article_no = m.group(3)

            # 往前找最近的 SEA。其後依序為：POL Country、POL Port、POD Country、POD Port。
            sea_idx: Optional[int] = None
            for j in range(i - 1, max(-1, i - 15), -1):
                if lines[j].upper() == "SEA":
                    sea_idx = j
                    break

            if sea_idx is None or sea_idx + 4 >= len(lines):
                continue

            pod_country = lines[sea_idx + 3]
            pod_port = lines[sea_idx + 4]
            pod = f"{pod_port}, {pod_country}"

            shipment_map[kal_article_no] = {
                "POD": pod,
                "LRD": lrd,
            }

    return shipment_map


def write_xlsx(
    items: List[OrderItem],
    validations: List[ValidationResult],
    output_path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Order Data"

    for col_idx, header in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(name="Arial", bold=True)

    # 檢核結果固定輸出至 L 欄，每張 PO 一列。
    ws["L1"] = "Validation Result"
    ws["L1"].font = Font(name="Arial", bold=True)

    for row_idx, item in enumerate(items, start=2):
        # LRD 原始格式為 DD/MM/YYYY，例如 11/01/2027 = 2027/1/11。
        # 轉成真正的 datetime，讓 Excel 可正常排序、篩選與日期運算。
        lrd_value = None
        if item.lrd:
            try:
                lrd_value = datetime.strptime(item.lrd, "%d/%m/%Y")
            except ValueError as exc:
                raise ValueError(f"LRD 日期格式錯誤：{item.lrd}") from exc

        values = [
            item.po_number,
            item.banner_article_no,
            item.article_description,
            item.quantity,
            item.unit_cost,
            item.amount,
            item.pod,
            lrd_value,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name="Arial")

        ws.cell(row=row_idx, column=1).number_format = "@"
        ws.cell(row=row_idx, column=2).number_format = "@"
        ws.cell(row=row_idx, column=8).number_format = "yyyy/m/d"
        ws.cell(row=row_idx, column=5).number_format = "0.00"
        ws.cell(row=row_idx, column=6).number_format = "#,##0.00"

    for row_idx, result in enumerate(validations, start=2):
        if result.is_match:
            message = (
                f"PO#{result.po_number}, 總數量及總金額與訂單相符"
            )
        else:
            message = (
                f"PO#{result.po_number}, 檢核不符："
                f"PDF總數量 {result.pdf_total_qty:,} / 輸出總數量 {result.output_total_qty:,}；"
                f"PDF總金額 {result.pdf_total_amount:,.2f} / "
                f"輸出總金額 {result.output_total_amount:,.2f}"
            )

        cell = ws.cell(row=row_idx, column=12, value=message)
        cell.font = Font(name="Arial")

    widths = {
        1: 14,
        2: 20,
        3: 55,
        4: 12,
        5: 12,
        6: 14,
        7: 22,
        8: 14,
        12: 95,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    if items:
        ws.auto_filter.ref = f"A1:H{len(items) + 1}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def extract_order(pdf_path: Path) -> tuple[List[OrderItem], ValidationResult]:
    if not pdf_path.exists():
        raise FileNotFoundError(f"找不到 PDF：{pdf_path}")

    with fitz.open(pdf_path) as doc:
        full_text = "\n".join(page.get_text("text") for page in doc)
        po_number = extract_po_number(full_text)
        pdf_total_qty, pdf_total_amount = extract_pdf_total(full_text)

        items = extract_article_items(doc, po_number)
        shipment_map = extract_shipment_info(doc)

        for item in items:
            shipment = shipment_map.get(item.kal_article_no, {})
            item.pod = shipment.get("POD", "")
            item.lrd = shipment.get("LRD", "")

    output_total_qty = sum(item.quantity for item in items)
    output_total_amount = round(sum(item.amount for item in items), 2)

    validation = ValidationResult(
        po_number=po_number,
        pdf_total_qty=pdf_total_qty,
        pdf_total_amount=pdf_total_amount,
        output_total_qty=output_total_qty,
        output_total_amount=output_total_amount,
        is_match=(
            output_total_qty == pdf_total_qty
            and abs(output_total_amount - pdf_total_amount) < 0.005
        ),
    )

    return items, validation


def extract_orders(pdf_paths: List[Path], output_path: Path) -> tuple[List[OrderItem], List[ValidationResult]]:
    """批次解析多張訂單，合併輸出至同一份 Excel。"""
    all_items: List[OrderItem] = []
    validations: List[ValidationResult] = []

    for pdf_path in pdf_paths:
        items, validation = extract_order(pdf_path)
        all_items.extend(items)
        validations.append(validation)

    write_xlsx(all_items, validations, output_path)
    return all_items, validations


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract Kingfisher / Castorama PO PDFs to Excel.")
    parser.add_argument("paths", nargs="*", help="Source PDF path(s), followed by output XLSX path")
    return parser.parse_args()


def run_gui() -> int:
    """未帶參數執行時，提供多選 PDF 與單一 Excel 輸出介面。"""
    root = tk.Tk()
    root.withdraw()

    try:
        pdf_files = filedialog.askopenfilenames(
            title="Select Kingfisher / Castorama PO PDFs",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")]
        )
        if not pdf_files:
            return 0

        pdf_paths = [Path(file) for file in pdf_files]

        if len(pdf_paths) == 1:
            default_name = f"{pdf_paths[0].stem}_extracted.xlsx"
        else:
            default_name = "Castorama_orders_extracted.xlsx"

        output_file = filedialog.asksaveasfilename(
            title="Save extracted Excel file",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel Workbook", "*.xlsx")]
        )
        if not output_file:
            return 0

        output_path = Path(output_file)
        items, validations = extract_orders(pdf_paths, output_path)

        matched = sum(1 for result in validations if result.is_match)
        failed = len(validations) - matched

        messagebox.showinfo(
            "Completed",
            f"Extraction completed.\n\n"
            f"POs: {len(validations)}\n"
            f"Items: {len(items)}\n"
            f"Validation OK: {matched}\n"
            f"Validation Failed: {failed}\n\n"
            f"Output:\n{output_path}"
        )
        return 0

    except Exception as exc:
        messagebox.showerror("Error", str(exc))
        return 1
    finally:
        root.destroy()


def main() -> int:
    args = parse_args()

    # 直接執行 .py / EXE 且未帶參數時，開啟 GUI。
    if not args.paths:
        return run_gui()

    # CLI / VBA 模式：
    # python script.py input1.pdf input2.pdf ... output.xlsx
    if len(args.paths) < 2:
        print(
            "ERROR: Provide at least one PDF path followed by one output XLSX path.",
            file=sys.stderr,
        )
        return 2

    pdf_paths = [Path(path).expanduser().resolve() for path in args.paths[:-1]]
    output_path = Path(args.paths[-1]).expanduser().resolve()

    try:
        items, validations = extract_orders(pdf_paths, output_path)
        matched = sum(1 for result in validations if result.is_match)

        print(f"SUCCESS: {len(validations)} PO(s), {len(items)} item(s) extracted.")
        print(f"VALIDATION: {matched}/{len(validations)} PO(s) matched.")
        print(f"OUTPUT: {output_path}")

        if matched != len(validations):
            return 3
        return 0

    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
