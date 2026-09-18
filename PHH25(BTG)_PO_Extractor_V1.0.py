import os
import re
import threading
from datetime import datetime

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from tkinter import Tk, Frame, Label, Button, Listbox, Scrollbar, END, BOTH, LEFT, RIGHT, Y, X, filedialog, messagebox, StringVar
from tkinter import ttk


APP_TITLE = "PHH25 訂單資料擷取"
OUTPUT_PREFIX = "PHH25_PO_Data"
SHEET_NAME = "PO_Data"

HEADERS = [
    "File Name",
    "BTG PO#",
    "Line#",
    "Delivery Date",
    "BTG SKU#",
    "BTG DESCRIPTION",
    "BTG QTY ",
    "BTG UM",
]

# PDF 表格欄位的 X 座標區間。
# 依使用者提供的 PHH25 / Blinds To Go PO 格式設定。
COLUMN_RANGES = {
    "Line#": (0, 75),
    "BTG SKU#": (75, 165),
    "BTG DESCRIPTION": (165, 365),
    "Delivery Date": (350, 475),
    "BTG QTY ": (475, 540),
    "BTG UM": (540, 575),
}


def normalize_spaces(value):
    """將連續空白整理為單一空白。"""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_number(value):
    """可轉成數值則轉換；否則保留原文字。"""
    text = normalize_spaces(value).replace(",", "")
    if not text:
        return ""

    try:
        number = float(text)
        if number.is_integer():
            return int(number)
        return number
    except ValueError:
        return normalize_spaces(value)


def normalize_delivery_date(value):
    """
    將 PDF 可能拆成 '2026-07-2 0' 的日期還原為 '2026-07-20'。
    無法辨識時保留原文字，不因單一欄位異常中斷。
    """
    text = normalize_spaces(value)
    if not text:
        return ""

    compact = re.sub(r"\s+", "", text)

    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", compact)
    if match:
        try:
            return datetime.strptime(match.group(0), "%Y-%m-%d")
        except ValueError:
            return match.group(0)

    return text


def extract_po_number(page_text):
    """由頁首擷取 PO#；保留前導 0。"""
    if not page_text:
        return ""

    patterns = [
        r"\bPO\s*#?\s*(\d+)\s+PAGE\b",
        r"\bPO\s*#?\s*(\d+)\b",
    ]

    for pattern in patterns:
        match = re.search(pattern, page_text, flags=re.IGNORECASE)
        if match:
            return match.group(1)

    return ""


def group_words_by_line(words, tolerance=2.5):
    """
    依 PDF word 的 top 座標分組成視覺上的同一列。
    此方式比單純解析 extract_text() 更能保留欄位位置。
    """
    if not words:
        return []

    sorted_words = sorted(words, key=lambda w: (w["top"], w["x0"]))
    lines = []

    for word in sorted_words:
        if not lines or abs(word["top"] - lines[-1]["top"]) > tolerance:
            lines.append({"top": word["top"], "words": [word]})
        else:
            lines[-1]["words"].append(word)

    for line in lines:
        line["words"].sort(key=lambda w: w["x0"])

    return lines


def text_in_x_range(words, x_min, x_max):
    """取得指定 X 座標範圍內的文字。"""
    selected = [
        w["text"]
        for w in words
        if x_min <= ((w["x0"] + w["x1"]) / 2) < x_max
    ]
    return normalize_spaces(" ".join(selected))


def identify_item_line(words):
    """判斷該視覺列是否為訂單明細列，並取得 Line#。"""
    for word in words:
        center_x = (word["x0"] + word["x1"]) / 2
        if center_x >= COLUMN_RANGES["Line#"][1]:
            continue

        # 明細 Line# 在此格式中固定顯示為 1.00、2.00...
        # 必須要求 .00，避免將頁首地址中的 901 等數字誤判為訂單列。
        match = re.fullmatch(r"(\d+)\.00", word["text"].strip())
        if match:
            return int(match.group(1))

    return None


def find_delivery_date(lines, item_index):
    """
    Delivery Date 位於 item 明細列的下一視覺列。
    僅向下查看少量列，避免誤抓下一筆資料或頁尾文字。
    """
    item_top = lines[item_index]["top"]

    for next_index in range(item_index + 1, min(item_index + 4, len(lines))):
        next_line = lines[next_index]
        vertical_gap = next_line["top"] - item_top

        if vertical_gap <= 0:
            continue
        if vertical_gap > 22:
            break

        # 若已遇到下一筆 item，不再往下找。
        if identify_item_line(next_line["words"]) is not None:
            break

        raw_date = text_in_x_range(
            next_line["words"],
            *COLUMN_RANGES["Delivery Date"]
        )

        if raw_date:
            return normalize_delivery_date(raw_date)

    return ""


def extract_page_rows(page, file_name):
    """
    擷取單頁訂單資料。
    各欄位獨立讀取：缺什麼就留空白，不要求整列欄位必須完整。
    """
    page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
    po_number = extract_po_number(page_text)

    words = page.extract_words(
        x_tolerance=2,
        y_tolerance=3,
        keep_blank_chars=False,
        use_text_flow=False,
    )

    lines = group_words_by_line(words)
    rows = []

    for index, line in enumerate(lines):
        line_number = identify_item_line(line["words"])
        if line_number is None:
            continue

        sku = text_in_x_range(
            line["words"],
            *COLUMN_RANGES["BTG SKU#"]
        )

        description = text_in_x_range(
            line["words"],
            *COLUMN_RANGES["BTG DESCRIPTION"]
        )

        qty = text_in_x_range(
            line["words"],
            *COLUMN_RANGES["BTG QTY "]
        )

        um = text_in_x_range(
            line["words"],
            *COLUMN_RANGES["BTG UM"]
        )
        if um:
            # UM 欄只取第一個單位文字，避免第 1 筆右側 COST 說明被一起抓入。
            um = um.split()[0]

        delivery_date = find_delivery_date(lines, index)

        rows.append({
            "File Name": file_name,
            "BTG PO#": po_number,
            "Line#": line_number,
            "Delivery Date": delivery_date,
            "BTG SKU#": parse_number(sku) if sku else "",
            "BTG DESCRIPTION": description,
            "BTG QTY ": parse_number(qty) if qty else "",
            "BTG UM": um,
        })

    return rows


def extract_pdf_rows(pdf_path):
    """擷取單一 PDF 全部頁面的訂單資料。"""
    rows = []
    file_name = os.path.basename(pdf_path)

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            rows.extend(extract_page_rows(page, file_name))

    return rows


def create_excel(rows, output_path):
    """依 PHH25.xlsx 範例欄位輸出 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    for col_index, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin = Side(style="thin", color="B7B7B7")

    for row_index, row_data in enumerate(rows, start=2):
        for col_index, header in enumerate(HEADERS, start=1):
            value = row_data.get(header, "")
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.font = Font(name="Arial")
            cell.border = Border(bottom=thin)

            if header == "BTG PO#":
                cell.number_format = "@"

            elif header == "Delivery Date" and isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd"

    # 欄寬依內容自動調整，並限制最大寬度避免 Description 過寬。
    for col_index, header in enumerate(HEADERS, start=1):
        max_length = len(header)

        for row_index in range(2, ws.max_row + 1):
            value = ws.cell(row=row_index, column=col_index).value
            if value is not None:
                max_length = max(max_length, len(str(value)))

        width = min(max(max_length + 2, 10), 45)
        ws.column_dimensions[get_column_letter(col_index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"

    wb.save(output_path)


class OrderExtractorApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("760x520")
        self.root.minsize(680, 460)

        self.pdf_files = []
        self.output_folder = ""

        self.status_var = StringVar(value="請先選擇訂單 PDF。")
        self.output_var = StringVar(value="尚未選擇輸出資料夾")

        self.build_ui()

    def build_ui(self):
        main = Frame(self.root, padx=16, pady=16)
        main.pack(fill=BOTH, expand=True)

        Label(main, text="訂單來源 PDF", font=("Arial", 11, "bold")).pack(anchor="w")

        file_frame = Frame(main)
        file_frame.pack(fill=BOTH, expand=True, pady=(6, 10))

        scrollbar = Scrollbar(file_frame)
        scrollbar.pack(side=RIGHT, fill=Y)

        self.file_list = Listbox(
            file_frame,
            yscrollcommand=scrollbar.set,
            font=("Arial", 10),
            selectmode="extended",
        )
        self.file_list.pack(side=LEFT, fill=BOTH, expand=True)
        scrollbar.config(command=self.file_list.yview)

        button_frame = Frame(main)
        button_frame.pack(fill=X, pady=(0, 12))

        Button(
            button_frame,
            text="選擇訂單 PDF（可多選）",
            command=self.select_files,
            width=22,
        ).pack(side=LEFT)

        Button(
            button_frame,
            text="清除清單",
            command=self.clear_files,
            width=12,
        ).pack(side=LEFT, padx=(8, 0))

        Label(main, text="輸出位置", font=("Arial", 11, "bold")).pack(anchor="w")

        output_frame = Frame(main)
        output_frame.pack(fill=X, pady=(6, 14))

        Label(
            output_frame,
            textvariable=self.output_var,
            anchor="w",
            relief="sunken",
            padx=8,
        ).pack(side=LEFT, fill=X, expand=True)

        Button(
            output_frame,
            text="選擇資料夾",
            command=self.select_output_folder,
            width=12,
        ).pack(side=RIGHT, padx=(8, 0))

        self.progress = ttk.Progressbar(main, mode="determinate", maximum=100)
        self.progress.pack(fill=X, pady=(0, 8))

        Label(
            main,
            textvariable=self.status_var,
            anchor="w",
            font=("Arial", 10),
        ).pack(fill=X, pady=(0, 10))

        self.start_button = Button(
            main,
            text="開始轉換",
            command=self.start_processing,
            height=2,
            font=("Arial", 11, "bold"),
        )
        self.start_button.pack(fill=X)

    def select_files(self):
        files = filedialog.askopenfilenames(
            title="選擇訂單 PDF",
            filetypes=[("PDF Files", "*.pdf")],
        )

        if not files:
            return

        existing = set(self.pdf_files)

        for file_path in files:
            if file_path not in existing:
                self.pdf_files.append(file_path)
                existing.add(file_path)

        self.refresh_file_list()
        self.status_var.set(f"已選擇 {len(self.pdf_files)} 個 PDF。")

    def clear_files(self):
        self.pdf_files.clear()
        self.refresh_file_list()
        self.progress["value"] = 0
        self.status_var.set("訂單清單已清除。")

    def refresh_file_list(self):
        self.file_list.delete(0, END)

        for file_path in self.pdf_files:
            self.file_list.insert(END, file_path)

    def select_output_folder(self):
        folder = filedialog.askdirectory(title="選擇轉出檔案存放位置")

        if not folder:
            return

        self.output_folder = folder
        self.output_var.set(folder)

    def start_processing(self):
        if not self.pdf_files:
            messagebox.showwarning("提醒", "請先選擇至少一個訂單 PDF。")
            return

        if not self.output_folder:
            messagebox.showwarning("提醒", "請先選擇轉出檔案存放位置。")
            return

        self.start_button.config(state="disabled")
        self.progress["value"] = 0
        self.status_var.set("準備開始處理...")

        thread = threading.Thread(target=self.process_files, daemon=True)
        thread.start()

    def process_files(self):
        all_rows = []
        errors = []
        total_files = len(self.pdf_files)

        try:
            for index, pdf_path in enumerate(self.pdf_files, start=1):
                file_name = os.path.basename(pdf_path)
                self.set_status(f"處理中：{file_name}")

                try:
                    rows = extract_pdf_rows(pdf_path)
                    all_rows.extend(rows)

                    if not rows:
                        errors.append(f"{file_name}：未辨識到訂單明細")
                except Exception as exc:
                    errors.append(f"{file_name}：{exc}")

                progress_value = (index / total_files) * 90
                self.set_progress(progress_value)

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_name = f"{OUTPUT_PREFIX}_{timestamp}.xlsx"
            output_path = os.path.join(self.output_folder, output_name)

            self.set_status("正在產生 Excel...")
            create_excel(all_rows, output_path)
            self.set_progress(100)

            summary = (
                f"轉換完成。\n\n"
                f"PDF 數量：{total_files}\n"
                f"輸出明細：{len(all_rows)} 筆\n"
                f"輸出檔案：\n{output_path}"
            )

            if errors:
                summary += "\n\n注意事項：\n" + "\n".join(errors)

            self.root.after(
                0,
                lambda: messagebox.showinfo("完成", summary),
            )

            self.set_status(f"完成：共輸出 {len(all_rows)} 筆訂單明細。")

        except Exception as exc:
            error_message = f"轉換失敗：{exc}"
            self.set_status(error_message)
            self.root.after(
                0,
                lambda msg=error_message: messagebox.showerror("錯誤", msg),
            )

        finally:
            self.root.after(
                0,
                lambda: self.start_button.config(state="normal"),
            )

    def set_status(self, text):
        self.root.after(0, lambda: self.status_var.set(text))

    def set_progress(self, value):
        self.root.after(0, lambda: self.progress.config(value=value))


def main():
    root = Tk()
    OrderExtractorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
